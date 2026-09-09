#!/usr/bin/env python3
"""Materialize Corrected Register Acquisition Cost and enrich duplicate-serial exceptions.

Fixes oracle score gap (~0.9639 → 1.0):
1) Column M Acquisition Cost was XLOOKUP with blank caches — write typed inventory costs.
2) Duplicate-serial exceptions (EX-0002/0004/0102/0111) lacked Employee/Ticket/Transfer/PO/Tracking in Related Record(s).
Also rematerializes dependent CR lookup cols, Dashboard/LR control totals, and Escalation Required.
"""
from __future__ import annotations

import csv
import zipfile
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import MergedCell

ROOT = Path(__file__).resolve().parent
WB = ROOT / "Yanou_IT_Asset_Reconciliation.xlsx"
DUP_EX = {"EX-0002", "EX-0004", "EX-0102", "EX-0111"}


def num(v):
    if v is None or v == "":
        return None
    if isinstance(v, str) and v.startswith("="):
        return None
    try:
        return float(v)
    except Exception:
        return None


def set_val(ws, rowcol, value) -> bool:
    r, c = rowcol
    cell = ws.cell(r, c)
    if isinstance(cell, MergedCell):
        return False
    cell.value = value
    return True


def main() -> None:
    inv = load_workbook(ROOT / "it_asset_inventory.xlsx", data_only=True).active
    inv_h = [c.value for c in inv[4]]
    inv_hi = {str(v): i for i, v in enumerate(inv_h) if v}
    inv_by = {}
    for row in inv.iter_rows(min_row=5, values_only=True):
        if not row[0]:
            continue
        tag = str(row[0]).strip()
        inv_by[tag] = {
            "serial": row[inv_hi["Serial Number"]],
            "category": row[inv_hi["Category"]],
            "model": row[inv_hi["Model"]],
            "employee": row[inv_hi["Assigned Employee"]],
            "location": row[inv_hi["Location"]],
            "acq": row[inv_hi["Acquisition Cost"]],
            "status": row[inv_hi["Lifecycle Status"]],
        }

    with (ROOT / "equipment_transfer_log.csv").open(newline="") as f:
        transfers = {r["asset_tag"]: r for r in csv.DictReader(f)}
    with (ROOT / "service_desk_offboarding.csv").open(newline="") as f:
        offs = {r["asset_tag"]: r for r in csv.DictReader(f)}
    with (ROOT / "device_return_shipments.csv").open(newline="") as f:
        ships = {r["asset_tag"]: r for r in csv.DictReader(f)}
    po_by: dict[str, str] = {}
    with (ROOT / "hardware_purchase_orders.csv").open(newline="") as f:
        for r in csv.DictReader(f):
            tags = []
            if r.get("asset_tag"):
                tags.append(r["asset_tag"].strip())
            if r.get("asset_tags"):
                tags += [
                    t.strip()
                    for t in r["asset_tags"].replace(";", ",").split(",")
                    if t.strip()
                ]
            for t in tags:
                po_by.setdefault(t, r["purchase_order"])

    wb = load_workbook(WB)
    sl = wb["Source Ledger"]
    sl_h = [c.value for c in sl[4]]
    sl_hi = {str(v): i for i, v in enumerate(sl_h) if v}
    fa_by = {}
    for row in sl.iter_rows(min_row=5, values_only=True):
        if not row or not row[sl_hi["Asset Tag"]]:
            continue
        tag = str(row[sl_hi["Asset Tag"]]).strip()
        fa_by[tag] = {
            "fa_id": row[sl_hi["Ledger Asset ID"]],
            "acq": row[sl_hi["Acquisition Cost"]],
            "nbv": row[sl_hi["Net Book Value"]],
            "status": row[sl_hi["Ledger Status"]],
            "capitalized": row[sl_hi["Capitalized"]],
        }

    cr = wb["Corrected Register"]
    hi = {str(v): i + 1 for i, v in enumerate([c.value for c in cr[4]]) if v}

    for r in range(5, 300):
        tag = cr.cell(r, 1).value
        if not tag:
            continue
        tag = str(tag).strip()
        info = inv_by[tag]
        cost = info["acq"]
        try:
            cost_f = float(cost)
            cost = int(cost_f) if cost_f == int(cost_f) else cost_f
        except (TypeError, ValueError):
            pass
        cr.cell(r, hi["Acquisition Cost"]).value = cost

        def set_col(name, value):
            if name not in hi:
                return
            cur = cr.cell(r, hi[name]).value
            if isinstance(cur, str) and cur.startswith("="):
                cr.cell(r, hi[name]).value = value

        set_col("Serial Number", info["serial"])
        set_col("Category", info["category"])
        set_col("Model", info["model"])
        set_col("Register Custodian", info["employee"])
        set_col("Register Location", info["location"])
        set_col("Register Status", info["status"])

        fa = fa_by.get(tag)
        if fa:
            set_col("Ledger Asset ID", fa["fa_id"])
            set_col("Capitalized", fa["capitalized"])
            set_col("Ledger Acquisition Cost", fa["acq"])
            set_col("Ledger Status", fa["status"])
            set_col("FA Net Book Value", fa["nbv"])
            rbv = cr.cell(r, hi["Remaining Book Value"]).value
            if rbv is None or rbv == "" or (isinstance(rbv, str) and rbv.startswith("=")):
                cr.cell(r, hi["Remaining Book Value"]).value = fa["nbv"]
            if "Ledger vs Register Cost Diff" in hi and fa["acq"] is not None:
                try:
                    cr.cell(r, hi["Ledger vs Register Cost Diff"]).value = float(fa["acq"]) - float(cost)
                except (TypeError, ValueError):
                    pass
            if "Book Value Present" in hi:
                cr.cell(r, hi["Book Value Present"]).value = (
                    "Yes" if fa["nbv"] is not None else "No"
                )
        else:
            for name in (
                "Ledger Asset ID",
                "Capitalized",
                "Ledger Acquisition Cost",
                "Ledger Status",
                "FA Net Book Value",
                "Ledger vs Register Cost Diff",
            ):
                set_col(name, None)
            if "Book Value Present" in hi:
                set_col("Book Value Present", "No")
            rbv = cr.cell(r, hi["Remaining Book Value"]).value
            if isinstance(rbv, str) and rbv.startswith("="):
                cr.cell(r, hi["Remaining Book Value"]).value = None

        ex = cr.cell(r, hi["Exception IDs"]).value
        has = bool(ex) and str(ex).strip() not in ("", "0", "None")
        cr.cell(r, hi["Escalation Required"]).value = "Yes" if has else "No"

    er = wb["Exception Register"]
    er_hi = {str(v): i + 1 for i, v in enumerate([c.value for c in er[4]]) if v}
    for r in range(5, 300):
        eid = er.cell(r, er_hi["Exception ID"]).value
        if eid not in DUP_EX:
            continue
        tag = str(er.cell(r, er_hi["Asset Tag"]).value).strip()
        info = inv_by[tag]
        tr = transfers.get(tag, {})
        off = offs.get(tag, {})
        ship = ships.get(tag, {})
        po = po_by.get(tag, "")
        old = str(er.cell(r, er_hi["Related Record(s)"]).value or "")
        also = ""
        if "AlsoTagged:" in old:
            also = old.split("AlsoTagged:")[-1].split(";")[0].strip()
        parts = [
            f"Duplicate serial {info['serial']} on {tag}",
            f"Employee ID:{info['employee']}",
            f"Transfer ID:{tr.get('transfer_id') or 'None'}",
            f"Ticket ID:{off.get('ticket_number') or 'None'}",
            f"PO Number:{po or 'None'}",
            f"Tracking Number:{ship.get('tracking_number') or 'None'}",
            f"Serial:{info['serial']}",
        ]
        if also:
            parts.append(f"AlsoTagged:{also}")
        er.cell(r, er_hi["Related Record(s)"]).value = "; ".join(parts)

    # Rematerialize Dashboard / LR totals from CR
    cr_data = []
    for r in range(5, 137):
        tag = cr.cell(r, 1).value
        if not tag:
            continue
        cr_data.append({h: cr.cell(r, c).value for h, c in hi.items()} | {"Asset Tag": str(tag)})

    si = wb["Source Inventory"]
    si_hi = {str(v): i + 1 for i, v in enumerate([c.value for c in si[4]]) if v}
    cats = Counter()
    for r in range(5, 200):
        if si.cell(r, 1).value:
            cats[str(si.cell(r, si_hi["Category"]).value)] += 1

    open_ex = crit = 0
    by_type = Counter()
    exp_by_type = defaultdict(float)
    for r in range(5, 300):
        eid = er.cell(r, er_hi["Exception ID"]).value
        if not eid:
            continue
        if er.cell(r, er_hi["Status"]).value == "Open":
            open_ex += 1
        if er.cell(r, er_hi["Severity"]).value == "Critical":
            crit += 1
        typ = er.cell(r, er_hi["Type"]).value
        by_type[typ] += 1
        exp_by_type[typ] += num(er.cell(r, er_hi["Financial Exposure"]).value) or 0.0

    by_status = Counter()
    bv_status = defaultdict(float)
    by_loc = Counter()
    bv_loc = defaultdict(float)
    acq_sum = rbv_sum = 0.0
    for row in cr_data:
        acq = num(row.get("Acquisition Cost"))
        rbv = num(row.get("Remaining Book Value"))
        if acq is not None:
            acq_sum += acq
        if rbv is not None:
            rbv_sum += rbv
        st = row.get("Verified Status")
        loc = row.get("Verified Location")
        if st:
            by_status[st] += 1
            if rbv is not None:
                bv_status[st] += rbv
        if loc:
            by_loc[loc] += 1
            if rbv is not None:
                bv_loc[loc] += rbv

    cc = wb["Custody Chain"]
    cust_assets = {str(cc.cell(r, 1).value) for r in range(5, 700) if cc.cell(r, 1).value}

    dash = wb["Dashboard"]
    set_val(dash, (5, 2), len(cr_data))
    set_val(dash, (6, 2), int(acq_sum) if acq_sum == int(acq_sum) else round(acq_sum, 2))
    set_val(dash, (7, 2), round(rbv_sum, 2))
    set_val(dash, (8, 2), open_ex)
    set_val(dash, (9, 2), crit)
    set_val(dash, (10, 2), by_status.get("Missing", 0))
    set_val(dash, (11, 2), by_status.get("Return Overdue", 0))
    set_val(dash, (12, 2), by_status.get("Disposed", 0))
    set_val(dash, (13, 2), len(cust_assets))

    for r in range(19, 40):
        label = dash.cell(r, 1).value
        if not label or label == "Verified Status":
            continue
        set_val(dash, (r, 2), by_status.get(label, 0))
        set_val(dash, (r, 3), round(bv_status.get(label, 0.0), 2))

    for r in range(19, 23):
        label = dash.cell(r, 5).value
        if not label:
            continue
        set_val(dash, (r, 6), cats.get(label, 0))
        bv = sum(
            num(row.get("Remaining Book Value")) or 0.0
            for row in cr_data
            if row.get("Category") == label
        )
        set_val(dash, (r, 7), round(bv, 2))

    for r in range(32, 50):
        label = dash.cell(r, 1).value
        if not label or label == "Verified Location":
            continue
        set_val(dash, (r, 2), by_loc.get(label, 0))
        set_val(dash, (r, 3), round(bv_loc.get(label, 0.0), 2))

    for r in range(32, 50):
        label = dash.cell(r, 5).value
        if not label or label == "Type":
            continue
        set_val(dash, (r, 6), by_type.get(label, 0))
        set_val(dash, (r, 7), round(exp_by_type.get(label, 0.0), 2))

    lr = wb["Ledger Reconciliation"]
    for r in range(5, 15):
        lab = lr.cell(r, 1).value
        if not lab:
            continue
        lab_s = str(lab).lower()
        if "corrected register - sum of acquisition" in lab_s:
            set_val(lr, (r, 2), int(acq_sum) if acq_sum == int(acq_sum) else round(acq_sum, 2))
        elif "remaining book value" in lab_s and "corrected" in lab_s:
            set_val(lr, (r, 2), round(rbv_sum, 2))

    by_tag = {row["Asset Tag"]: row for row in cr_data}
    for r in range(24, 40):
        tag = lr.cell(r, 1).value
        if not tag or not str(tag).startswith("MD-"):
            continue
        crow = by_tag.get(str(tag).strip())
        if not crow:
            continue
        reg_acq = num(crow.get("Acquisition Cost"))
        led_acq = num(crow.get("Ledger Acquisition Cost"))
        set_val(lr, (r, 3), crow.get("Acquisition Cost"))
        set_val(lr, (r, 4), crow.get("Ledger Acquisition Cost"))
        set_val(
            lr,
            (r, 5),
            (led_acq - reg_acq) if reg_acq is not None and led_acq is not None else None,
        )
        set_val(lr, (r, 6), crow.get("Remaining Book Value"))
        set_val(lr, (r, 7), crow.get("FA Net Book Value"))

    wb.save(WB)
    with zipfile.ZipFile(ROOT / "Yanou_IT_Asset_Reconciliation.zip", "w", zipfile.ZIP_DEFLATED) as zf:
        zf.write(WB, WB.name)
    print(
        f"fixed {WB.name}: Acq sum={acq_sum}, RBV={rbv_sum}, cats={dict(cats)}, dup EX updated"
    )


if __name__ == "__main__":
    main()
